"""
Script: Concatenate.py
Objetivo: Concatena todos los Excel de la carpeta Excel en un unico df_master_01.xlsx
"""
import openpyxl
from pathlib import Path

EXCEL_DIR   = r"C:\Users\victor\Desktop\Nueva carpeta (2)\Excel"
OUTPUT_FILE = r"C:\Users\victor\Desktop\Nueva carpeta (2)\Excel\df_master_01.xlsx"

HEADERS = [
    "Year", "Region_Code", "Region_Name", "Province",
    "Location", "Level",
    "Foreign_Travelers", "OFW", "Domestic_Travelers", "Total",
]

def main():
    excel_dir = Path(EXCEL_DIR)

    # Recoger solo los excels de año (excluir df_master si ya existe)
    files = sorted(
        f for f in excel_dir.glob("Philippines_Tourism_*.xlsx")
        if f.name != "df_master_01.xlsx"
    )

    if not files:
        print("No se encontraron archivos Philippines_Tourism_*.xlsx")
        return

    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active
    ws_out.title = "df_master"
    ws_out.append(HEADERS)

    seen = set()  # clave unica para detectar duplicados
    total_rows = 0
    dupes = 0

    for f in files:
        wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        wb.close()

        for row in rows:
            if not any(row):
                continue
            # Clave: Year + Region_Code + Location + Level (identifica la fila de forma unica)
            key = (row[0], row[1], row[4], row[5])
            if key in seen:
                dupes += 1
                continue
            seen.add(key)
            ws_out.append(list(row))
            total_rows += 1

        print(f"  {f.name}: {len(rows)} filas leidas")

    # Ajustar anchos de columna
    col_widths = [6, 12, 45, 30, 35, 18, 18, 10, 18, 14]
    for i, w in enumerate(col_widths, 1):
        ws_out.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    wb_out.save(OUTPUT_FILE)
    print(f"\nFilas totales escritas : {total_rows}")
    print(f"Duplicados eliminados  : {dupes}")
    print(f"Guardado en            : {OUTPUT_FILE}")

if __name__ == "__main__":
    main()
